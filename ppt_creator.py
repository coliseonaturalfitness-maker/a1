#!/usr/bin/env python3
"""Creador simple de presentaciones PPTX desde JSON."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from pptx import Presentation


def build_presentation(input_path: Path, output_path: Path) -> None:
    data = json.loads(input_path.read_text(encoding="utf-8"))
    slides = data.get("slides", [])

    if not slides:
        raise ValueError("El archivo JSON debe incluir al menos una diapositiva en 'slides'.")

    prs = Presentation()

    for idx, slide_data in enumerate(slides):
        title = slide_data.get("title", f"Diapositiva {idx + 1}")
        bullets = slide_data.get("bullets", [])

        layout = prs.slide_layouts[1]  # Title and Content
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = title

        body = slide.shapes.placeholders[1].text_frame
        body.clear()

        if bullets:
            first = body.paragraphs[0]
            first.text = str(bullets[0])
            for bullet in bullets[1:]:
                p = body.add_paragraph()
                p.text = str(bullet)
                p.level = 0

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Genera una presentación .pptx a partir de un archivo JSON."
    )
    parser.add_argument("input", type=Path, help="Ruta del archivo JSON de entrada")
    parser.add_argument("output", type=Path, help="Ruta del archivo .pptx de salida")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    build_presentation(args.input, args.output)
    print(f"Presentación creada: {args.output}")


if __name__ == "__main__":
    main()
